"""Reproducible waste classification experiments. Run `python research.py --help`."""
import argparse
import hashlib
import html
import itertools
import json
import os
from pathlib import Path
import platform
import subprocess
import sys
import time
import urllib.request
import zipfile

import numpy as np
import pandas as pd
from PIL import Image, ImageFilter
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix, f1_score
from sklearn.model_selection import StratifiedGroupKFold

CLASSES = ['cardboard', 'glass', 'metal', 'paper', 'plastic', 'trash']
MODELS = ['custom_cnn', 'mobilenetv2', 'resnet50', 'efficientnetb0']
URL = 'https://raw.githubusercontent.com/garythung/trashnet/master/data/dataset-resized.zip'

def write_json(path, value):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, allow_nan=False), encoding='utf-8')

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def download(destination):
    destination = Path(destination)
    root = destination / 'dataset-resized'
    if root.exists():
        print('Dataset already exists:', root); return
    destination.mkdir(parents=True, exist_ok=True)
    archive = destination / 'trashnet.zip'
    print('Downloading TrashNet from the original repository ...')
    urllib.request.urlretrieve(URL, archive)
    with zipfile.ZipFile(archive) as z:
        for member in z.infolist():
            resolved = (destination / member.filename).resolve()
            if not resolved.is_relative_to(destination.resolve()):
                raise ValueError('Unsafe archive path')
        z.extractall(destination)
    write_json(destination / 'source.json', {'url': URL, 'archive_sha256': digest(archive),
        'downloaded_utc': time.strftime('%Y-%m-%dT%H:%M:%SZ', time.gmtime()),
        'citation': 'Gary Thung and Mindy Yang, TrashNet, https://github.com/garythung/trashnet'})
    print('Downloaded:', root)

def prepare(data, out):
    import imagehash
    data, out = Path(data).resolve(), Path(out)
    out.mkdir(parents=True, exist_ok=True)
    if (out / 'manifest.csv').exists():
        raise ValueError('Manifest already exists. Use another output directory to create a new split.')
    rows, excluded, seen, conflicts = [], [], {}, set()
    for label, category in enumerate(CLASSES):
        folder = data / category
        if not folder.is_dir(): raise ValueError(f'Missing class folder: {folder}')
        for path in sorted(folder.rglob('*')):
            if path.suffix.lower() not in ['.jpg', '.jpeg', '.png', '.bmp', '.webp']: continue
            try:
                with Image.open(path) as im:
                    im.load(); rgb = im.convert('RGB')
                    sha = hashlib.sha256(np.asarray(rgb).tobytes() + str(rgb.size).encode()).hexdigest()
                    ph = str(imagehash.phash(rgb))
                if sha in seen:
                    if seen[sha] != category:
                        conflicts.add(sha)
                        excluded.append({'path':str(path),'reason':'identical decoded image with conflicting labels; all copies excluded'})
                    else:
                        excluded.append({'path': str(path), 'reason': 'identical decoded image'})
                    continue
                seen[sha] = category
                rows.append({'path': path.relative_to(data).as_posix(), 'label': label, 'class': category,
                             'sha256': digest(path), 'phash': ph, 'decoded_sha':sha})
            except ValueError: raise
            except Exception as e: excluded.append({'path': str(path), 'reason': str(e)})
    kept=[]
    for row in rows:
        if row['decoded_sha'] in conflicts:
            excluded.append({'path':str(data/row['path']),'reason':'identical decoded image with conflicting labels; all copies excluded'})
        else:
            row.pop('decoded_sha');kept.append(row)
    rows=kept
    if not rows: raise ValueError('No readable images found')
    parents = list(range(len(rows)))
    def find(i):
        while parents[i] != i:
            parents[i] = parents[parents[i]]; i = parents[i]
        return i
    hashes = [int(r['phash'], 16) for r in rows]
    candidates = []
    # Conservative automatic grouping, including across labels, prevents near duplicates crossing splits.
    for i in range(len(rows)):
        for j in range(i):
            distance = (hashes[i] ^ hashes[j]).bit_count()
            if distance <= 4:
                parents[find(i)] = find(j)
                candidates.append({'image_a': rows[i]['path'], 'image_b': rows[j]['path'], 'phash_distance': distance})
    groups = [find(i) for i in range(len(rows))]
    labels = np.array([r['label'] for r in rows])
    # 20 stratified group folds: 14 train, 3 validation, 3 test; approximate 70/15/15.
    if min(np.bincount(labels, minlength=6)) < 20:
        raise ValueError('At least 20 usable images per class are needed for this splitting protocol.')
    splitter = StratifiedGroupKFold(n_splits=20, shuffle=True, random_state=42)
    for fold, (_, indices) in enumerate(splitter.split(np.zeros(len(rows)), labels, groups)):
        split = 'train' if fold < 14 else ('validation' if fold < 17 else 'test')
        for i in indices: rows[i].update(group=groups[i], split=split)
    frame = pd.DataFrame(rows)
    for split in ['train', 'validation', 'test']:
        if frame[frame.split == split].label.nunique() != 6:
            raise ValueError('Grouping produced a split missing a class; inspect duplicate groups and use another split design.')
    frame.to_csv(out / 'manifest.csv', index=False)
    pd.DataFrame(candidates, columns=['image_a','image_b','phash_distance']).to_csv(out / 'near_duplicates_to_review.csv', index=False)
    pd.DataFrame(excluded, columns=['path','reason']).to_csv(out / 'excluded_images.csv', index=False)
    counts = pd.crosstab(frame['class'], frame['split']).reindex(CLASSES)
    counts.to_csv(out / 'dataset_counts.csv')
    write_json(out / 'dataset.json', {'root': str(data), 'classes': CLASSES,
        'manifest_sha256': digest(out / 'manifest.csv'), 'split_method': '20 stratified group folds: 14/3/3',
        'near_duplicate_threshold': 4, 'n_images': len(frame), 'n_excluded': len(excluded)})
    print(counts); print('Review near_duplicates_to_review.csv BEFORE training. Rebuild in a new output folder if correcting labels or grouping.')

def environment(tf):
    return {'python': sys.version, 'platform': platform.platform(), 'processor': platform.processor(),
        'tensorflow': tf.__version__, 'devices': [str(x) for x in tf.config.list_physical_devices()],
        'packages': subprocess.check_output([sys.executable, '-m', 'pip', 'freeze'], text=True).splitlines(),
        'cpu_count': os.cpu_count(), 'precision': 'float32'}

def load_image(path, size=224):
    with Image.open(path) as im:
        return np.asarray(im.convert('RGB').resize((size, size), Image.Resampling.BILINEAR), dtype=np.float32)

def data_frame(out):
    meta = json.loads((out / 'dataset.json').read_text())
    frame = pd.read_csv(out / 'manifest.csv')
    if digest(out / 'manifest.csv') != meta['manifest_sha256']: raise ValueError('Manifest changed. Prepare a new experiment directory.')
    return meta, frame

def compile_model(tf, model, lr):
    model.compile(optimizer=tf.keras.optimizers.Adam(lr), loss='sparse_categorical_crossentropy', metrics=['accuracy'])

def build_model(tf, name, seed, weights='imagenet'):
    k = tf.keras; layers = k.layers
    inputs = k.Input((224, 224, 3), name='rgb_0_to_255')
    aug = k.Sequential([layers.RandomFlip('horizontal', seed=seed),
        layers.RandomRotation(.05, fill_mode='reflect', seed=seed+1),
        layers.RandomZoom(.1, fill_mode='reflect', seed=seed+2)], name='training_augmentation')
    x = aug(inputs)
    backbone = None
    if name == 'custom_cnn':
        x = layers.Rescaling(1/255)(x)
        for filters in [32, 64, 128]:
            x = layers.Conv2D(filters, 3, padding='same', activation='relu')(x)
            x = layers.MaxPooling2D()(x)
        x = layers.GlobalAveragePooling2D()(x)
        x = layers.Dense(128, activation='relu')(x)
    else:
        if name == 'mobilenetv2':
            x = layers.Rescaling(1/127.5, offset=-1)(x)
            constructor = k.applications.MobileNetV2
        elif name == 'resnet50':
            # Keras ResNet50 preprocessing: RGB -> BGR, ImageNet mean subtraction.
            x = k.ops.flip(x, axis=-1)
            x = x - k.ops.convert_to_tensor([103.939,116.779,123.68], dtype='float32')
            constructor = k.applications.ResNet50
        else:
            constructor = k.applications.EfficientNetB0  # built-in preprocessing expects 0..255
        backbone = constructor(include_top=False, weights=weights, input_shape=(224,224,3))
        backbone.trainable = False
        x = backbone(x, training=False)  # keep batch normalization statistics fixed during fine-tuning
        x = layers.GlobalAveragePooling2D()(x)
    x = layers.Dropout(.3, seed=seed+3)(x)
    model = k.Model(inputs, layers.Dense(6, activation='softmax')(x))
    compile_model(tf, model, .001)
    return model, backbone

def metric_summary(y, pred):
    report = classification_report(y, pred, labels=list(range(6)), target_names=CLASSES,
                                   output_dict=True, zero_division=0)
    return {'accuracy': float(accuracy_score(y,pred)), 'macro_precision': report['macro avg']['precision'],
        'macro_recall': report['macro avg']['recall'], 'macro_f1': report['macro avg']['f1-score'],
        'weighted_f1': report['weighted avg']['f1-score']}, report

def save_figures(run, history, y, pred):
    import matplotlib; matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    for field in ['accuracy', 'loss', 'macro_f1']:
        fig, ax = plt.subplots(figsize=(6,4))
        key = field if field != 'macro_f1' else 'val_macro_f1'
        if key in history: ax.plot(history[key], marker='.', label=key)
        if 'val_'+field in history and field != 'macro_f1': ax.plot(history['val_'+field], marker='.', label='validation')
        ax.set(xlabel='Epoch index (zero based)', ylabel=field, title=run.name); ax.legend(); fig.tight_layout()
        fig.savefig(run / f'{field}_curve.png', dpi=300); plt.close(fig)
    from sklearn.metrics import ConfusionMatrixDisplay
    fig, ax = plt.subplots(figsize=(7,6))
    ConfusionMatrixDisplay(confusion_matrix(y,pred,labels=range(6)), display_labels=CLASSES).plot(ax=ax, colorbar=False, xticks_rotation=45)
    fig.tight_layout(); fig.savefig(run / 'confusion_matrix.png',dpi=300); plt.close(fig)
    fig,axes=plt.subplots(1,2,figsize=(12,4))
    for ax,field in zip(axes,['accuracy','loss']):
        ax.plot(history[field],marker='.',label='training')
        ax.plot(history['val_'+field],marker='.',label='validation')
        ax.set(xlabel='Epoch index (zero based)',ylabel=field);ax.legend()
    fig.tight_layout();fig.savefig(run/'training_panel.png',dpi=300);plt.close(fig)

def corrupt(image, condition, index):
    if condition.startswith('brightness'): return np.clip(image * float(condition.split('_')[1]),0,255)
    if condition.startswith('blur'):
        return np.asarray(Image.fromarray(image.astype('uint8')).filter(ImageFilter.GaussianBlur(float(condition.split('_')[1]))),dtype=np.float32)
    rng=np.random.default_rng(2026+index)
    return np.clip(image+rng.normal(0,float(condition.split('_')[1])*255,image.shape),0,255).astype(np.float32)

def train(args):
    import tensorflow as tf
    out = Path(args.out); meta, frame = data_frame(out)
    root = Path(meta['root'])
    # Verify dataset bytes before training: prevents silent changes between runs.
    for r in frame.itertuples():
        if digest(root / r.path) != r.sha256: raise ValueError('Dataset image changed: '+r.path)
    settings = {'models': MODELS, 'seeds': args.seeds, 'batch_size': args.batch_size,
        'epochs': args.epochs, 'head_epochs': args.head_epochs, 'fine_epochs': args.fine_epochs,
        'weights': args.weights, 'smoke': args.smoke, 'manifest': meta['manifest_sha256'],
        'augmentation': {'rotation':.05,'zoom':.1,'flip':'horizontal'}, 'unfreeze_last_layers':30,
        'source_sha256':digest(__file__)}
    config_path=out/'experiment_config.json'
    if config_path.exists() and json.loads(config_path.read_text()) != settings:
        raise ValueError('Configuration differs from existing experiment. Use a new --out directory and prepare it first.')
    write_json(config_path,settings); write_json(out/'training_environment.json',environment(tf))
    def dataset(part, seed, shuffle=False):
        paths = [str(root / p) for p in part.path]
        ds=tf.data.Dataset.from_tensor_slices((paths,part.label.to_numpy(dtype='int32')))
        if shuffle: ds=ds.shuffle(len(part),seed=seed,reshuffle_each_iteration=True)
        def read(path,label):
            im=tf.numpy_function(lambda p:load_image(p.decode()),[path],tf.float32)
            im.set_shape((224,224,3));return im,label
        return ds.map(read,num_parallel_calls=tf.data.AUTOTUNE).batch(args.batch_size).prefetch(tf.data.AUTOTUNE)
    for name in args.models:
        for seed in args.seeds:
            run=out/'runs'/f'{name}_seed{seed}';run.mkdir(parents=True,exist_ok=True)
            if (run/'metrics.json').exists(): print('Skipping completed run:',run.name);continue
            tf.keras.backend.clear_session();tf.keras.utils.set_random_seed(seed)
            try:tf.config.experimental.enable_op_determinism()
            except Exception:pass
            tr=frame[frame.split=='train'];va=frame[frame.split=='validation'];te=frame[frame.split=='test']
            if args.smoke:
                tr=tr.groupby('label',group_keys=False).head(2);va=va.groupby('label',group_keys=False).head(1);te=te.groupby('label',group_keys=False).head(1)
            train_ds=dataset(tr,seed,True);val_ds=dataset(va,seed);test_ds=dataset(te,seed)
            yval=va.label.to_numpy();y=te.label.to_numpy()
            model,base=build_model(tf,name,seed,None if args.weights=='none' else 'imagenet')
            class_counts=np.bincount(tr.label,minlength=6);weights={i:len(tr)/(6*int(n)) for i,n in enumerate(class_counts)}
            class ValidationF1(tf.keras.callbacks.Callback):
                def on_epoch_end(self,epoch,logs=None):
                    pred=np.argmax(self.model.predict(val_ds,verbose=0),axis=1)
                    logs['val_macro_f1']=float(f1_score(yval,pred,labels=range(6),average='macro',zero_division=0))
                    print(f' validation macro F1: {logs["val_macro_f1"]:.4f}')
            class GlobalBest(tf.keras.callbacks.Callback):
                def __init__(self): super().__init__();self.best=-1
                def on_epoch_end(self,epoch,logs=None):
                    if logs['val_macro_f1']>self.best:
                        self.best=logs['val_macro_f1'];self.model.save_weights(run/'best.weights.h5')
            best=GlobalBest();hist={};st=time.perf_counter()
            stages=[('scratch',args.epochs)] if base is None else [('head',args.head_epochs),('fine',args.fine_epochs)]
            stage_records=[]
            for stage,epochs in stages:
                if epochs<=0:continue
                if stage=='fine':
                    model.load_weights(run/'best.weights.h5')
                    base.trainable=True
                    for layer in base.layers:layer.trainable=False
                    for layer in base.layers[-30:]:
                        if not isinstance(layer,tf.keras.layers.BatchNormalization):layer.trainable=True
                    compile_model(tf,model,.00001)
                h=model.fit(train_ds,validation_data=val_ds,epochs=epochs,class_weight=weights,
                    callbacks=[ValidationF1(),best,tf.keras.callbacks.EarlyStopping(monitor='val_macro_f1',mode='max',patience=5,restore_best_weights=False),
                               tf.keras.callbacks.CSVLogger(str(run/f'{stage}_training.csv'))],verbose=2)
                stage_records.append({'stage':stage,'epochs':len(h.history['loss']),
                    'trainable_parameters':int(sum(np.prod(w.shape) for w in model.trainable_weights))})
                for key,values in h.history.items():hist.setdefault(key,[]).extend(float(v) for v in values)
            duration=time.perf_counter()-st
            model.load_weights(run/'best.weights.h5')
            inference_model=tf.keras.Model(inputs=model.inputs,outputs=model.outputs,name=name+'_inference')
            inference_model.save(run/'model.keras')
            probs=model.predict(test_ds,verbose=0);pred=probs.argmax(axis=1)
            summary,rep=metric_summary(y,pred)
            write_json(run/'classification_report.json',rep)
            pd.DataFrame(rep).T.to_csv(run/'class_metrics.csv')
            pd.DataFrame(confusion_matrix(y,pred,labels=range(6)),index=CLASSES,columns=CLASSES).to_csv(run/'confusion_matrix.csv')
            predictions=te[['path','label','class']].copy();predictions['prediction']=pred
            for i,c in enumerate(CLASSES):predictions['prob_'+c]=probs[:,i]
            predictions.to_csv(run/'predictions.csv',index=False)
            robustness=[]
            if not args.smoke:
                conditions=['brightness_0.7','brightness_1.3','blur_1','blur_2','noise_0.02','noise_0.05']
                for condition in conditions:
                    changed=[]
                    paths=te.path.tolist()
                    for start in range(0,len(paths),args.batch_size):
                        images=np.stack([corrupt(load_image(root/p),condition,start+i) for i,p in enumerate(paths[start:start+args.batch_size])])
                        changed.extend(np.argmax(model(images,training=False).numpy(),axis=1).tolist())
                    cs,_=metric_summary(y,changed)
                    robustness.append({'condition':condition,'macro_f1':cs['macro_f1'], 'drop':summary['macro_f1']-cs['macro_f1']})
                pd.DataFrame(robustness).to_csv(run/'robustness.csv',index=False)
            write_json(run/'history.json',hist);save_figures(run,hist,y,pred)
            summary.update(model=name,seed=seed,smoke=args.smoke,weights=args.weights,n_test=len(te),
                training_seconds=duration,epochs=sum(s['epochs'] for s in stage_records),stages=stage_records,
                parameters=int(model.count_params()),model_size_mib=(run/'model.keras').stat().st_size/(1024**2))
            write_json(run/'metrics.json',summary) # completion marker, written last
            print('Completed',run.name)

def benchmark(args):
    # CPU environment is set before TensorFlow import; use a separate invocation after training.
    os.environ['CUDA_VISIBLE_DEVICES']='-1'
    import tensorflow as tf
    tf.config.set_visible_devices([], 'GPU')
    tf.config.threading.set_intra_op_parallelism_threads(args.threads)
    tf.config.threading.set_inter_op_parallelism_threads(1)
    out=Path(args.out);meta,frame=data_frame(out);paths=[Path(meta['root'])/p for p in frame[frame.split=='test'].path]
    write_json(out/'benchmark_environment.json',dict(environment(tf),threads=args.threads,calls=args.calls,warmup=20))
    for run in sorted((out/'runs').glob('*')):
        if not (run/'metrics.json').exists():continue
        tf.keras.backend.clear_session();model=tf.keras.models.load_model(run/'model.keras',compile=False)
        @tf.function(input_signature=[tf.TensorSpec([1,224,224,3],tf.float32)])
        def infer(x):return model(x,training=False)
        x=tf.constant(load_image(paths[0])[None])
        for _ in range(20):infer(x).numpy()
        raw=[];end=[]
        for i in range(args.calls):
            x=tf.constant(load_image(paths[i%len(paths)])[None])
            start=time.perf_counter_ns();infer(x).numpy();raw.append((time.perf_counter_ns()-start)/1e6)
            start=time.perf_counter_ns();im=load_image(paths[i%len(paths)]);infer(tf.constant(im[None])).numpy();end.append((time.perf_counter_ns()-start)/1e6)
        pd.DataFrame({'model_ms':raw,'end_to_end_ms':end}).to_csv(run/'latency_samples.csv',index=False)
        write_json(run/'latency.json',{'model_median_ms':float(np.median(raw)),'model_p95_ms':float(np.percentile(raw,95)),
            'end_to_end_median_ms':float(np.median(end)),'end_to_end_p95_ms':float(np.percentile(end,95)),
            'device':'CPU','threads':args.threads,'calls':args.calls,'float_precision':'float32','batch_size':1})
        print('Benchmarked',run.name)

def report(args):
    import matplotlib;matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    out=Path(args.out);paper=out/'paper_ready';paper.mkdir(exist_ok=True)
    records=[];run_dirs=[]
    for run in sorted((out/'runs').glob('*')):
        if not (run/'metrics.json').exists():continue
        r=json.loads((run/'metrics.json').read_text())
        if (run/'latency.json').exists():r.update(json.loads((run/'latency.json').read_text()))
        records.append(r);run_dirs.append(run)
    if not records:raise ValueError('No completed runs. Train first.')
    frame=pd.DataFrame(records);frame.drop(columns=['stages']).to_csv(paper/'all_runs.csv',index=False)
    metrics=['accuracy','macro_precision','macro_recall','macro_f1','weighted_f1','training_seconds','model_size_mib','parameters']
    metrics += [x for x in ['model_median_ms','model_p95_ms','end_to_end_median_ms'] if x in frame]
    aggregate=frame.groupby('model')[metrics].agg(['mean','std'])
    aggregate.columns=['_'.join(c) for c in aggregate.columns];aggregate.to_csv(paper/'comparison_mean_std.csv')
    clean=not frame.smoke.any() and (frame.weights!='none').all()
    complete=clean and set(frame.model)==set(MODELS) and all(set(frame[frame.model==m].seed)=={42,123,2026} for m in MODELS)
    title='Measured Waste Classification Results' if complete else 'Preliminary Waste Classification Results'
    sections=[f'# {title}', 'These results were generated from this experiment. Inspect them and their limitations before adding them to a manuscript.',
        '## Experimental status',f'Completed runs: {len(frame)}. Full standard 4-model × 3-seed study: {complete}. CPU latency available for all runs: {all((r/"latency.json").exists() for r in run_dirs)}.',
        '## Model comparison',aggregate.round(5).to_csv(),
        'Values in CSV tables are fractions for accuracy and F1 (multiply by 100 for percentages). Standard deviation describes training-seed variability on one fixed split, not uncertainty across datasets. Model size is the .keras inference archive including preprocessing and weights. Training time includes validation prediction for macro F1 and checkpoint writing.']
    best_model=aggregate['macro_f1_mean'].idxmax();score=aggregate.loc[best_model,'macro_f1_mean']
    sections += ['## Results paragraph to edit', f'Among the completed configurations, {best_model} obtained the highest mean test macro F1 ({score:.4f}). This ranking is specific to the fixed dataset split and training protocol. Class-wise metrics and paired uncertainty analysis should be considered before interpreting a small difference as meaningful. No practical sorting reliability is established by this result alone.',
        '## Figures and tables', 'comparison_macro_f1.png: Mean test macro F1 with one training-seed standard deviation.\nperformance_vs_latency.png: CPU median inference latency versus mean macro F1, when timing is available.\nclass_metrics_all.csv: Per-class precision, recall and F1 for every run.\nrobustness_summary.csv: Mean and standard deviation under predeclared image corruptions.\nbootstrap_differences.csv: Exploratory paired stratified test-set confidence intervals.\nEach run folder contains accuracy/loss curves and a confusion matrix.',
        '## Limitations paragraph to edit','The evaluation uses a small dataset with controlled backgrounds and an imbalanced trash category. Three seeds share one split. Controlled corruptions simulate image degradation and do not demonstrate cross-dataset generalization. Inference timings apply only to the recorded CPU and software configuration. The scratch baseline and pretrained networks differ in architecture as well as initialization, so this experiment does not isolate a causal transfer-learning effect.']
    fig,ax=plt.subplots(figsize=(7,4))
    std=aggregate['macro_f1_std'].fillna(0)
    ax.bar(aggregate.index,aggregate.macro_f1_mean,yerr=std,capsize=4)
    ax.set(ylabel='Macro F1 (fraction)',ylim=(0,1),title=title);ax.tick_params(axis='x',rotation=15)
    fig.tight_layout();fig.savefig(paper/'comparison_macro_f1.png',dpi=300);plt.close(fig)
    if 'model_median_ms_mean' in aggregate:
        fig,ax=plt.subplots(figsize=(7,4))
        for m,r in aggregate.iterrows():
            ax.scatter(r.model_median_ms_mean,r.macro_f1_mean,label=m)
        ax.legend()
        ax.set(xlabel='Mean CPU batch-one median latency (ms)',ylabel='Mean macro F1',title='Performance and inference cost')
        fig.tight_layout();fig.savefig(paper/'performance_vs_latency.png',dpi=300);plt.close(fig)
        frontier=[]
        for m,row in aggregate.iterrows():
            dominated=False
            for other,r in aggregate.iterrows():
                if other==m:continue
                at_least=(r.macro_f1_mean>=row.macro_f1_mean and r.model_median_ms_mean<=row.model_median_ms_mean and r.model_size_mib_mean<=row.model_size_mib_mean)
                strict=(r.macro_f1_mean>row.macro_f1_mean or r.model_median_ms_mean<row.model_median_ms_mean or r.model_size_mib_mean<row.model_size_mib_mean)
                if at_least and strict:dominated=True
            frontier.append({'model':m,'non_dominated':not dominated})
        pd.DataFrame(frontier).to_csv(paper/'pareto_frontier.csv',index=False)
    perclass=[];rob=[]
    for run,r in zip(run_dirs,records):
        cr=json.loads((run/'classification_report.json').read_text())
        for c in CLASSES:perclass.append(dict(model=r['model'],seed=r['seed'],category=c,**cr[c]))
        if (run/'robustness.csv').exists():
            rb=pd.read_csv(run/'robustness.csv');rb['model']=r['model'];rb['seed']=r['seed'];rob.append(rb)
    pd.DataFrame(perclass).to_csv(paper/'class_metrics_all.csv',index=False)
    if rob:
        pd.concat(rob).groupby(['model','condition'])[['macro_f1','drop']].agg(['mean','std']).to_csv(paper/'robustness_summary.csv')
    # Paired stratified bootstrap of seed-averaged macro F1; preserves identical test-image ordering.
    pairs=[];rng=np.random.default_rng(2026)
    predictions={}
    for run,r in zip(run_dirs,records):predictions.setdefault(r['model'],[]).append(pd.read_csv(run/'predictions.csv'))
    for a,b in itertools.combinations(sorted(predictions),2):
        aa,bb=predictions[a],predictions[b]
        reference=aa[0]
        if any(not reference.path.equals(x.path) or not reference.label.equals(x.label) for x in aa+bb):
            raise ValueError('Test predictions have different image order or labels; cannot compare.')
        if {r['seed'] for r in records if r['model']==a}!={r['seed'] for r in records if r['model']==b}:
            continue
        y=reference.label.to_numpy(); strata=[np.flatnonzero(y==c) for c in range(6)]
        delta=[]
        for _ in range(args.bootstrap):
            ix=np.concatenate([rng.choice(s,len(s),replace=True) for s in strata if len(s)])
            def avg(rows):return np.mean([f1_score(y[ix],row.prediction.to_numpy()[ix],labels=range(6),average='macro',zero_division=0) for row in rows])
            delta.append(avg(aa)-avg(bb))
        pairs.append({'model_a':a,'model_b':b,'mean_f1_difference':float(aggregate.loc[a,'macro_f1_mean']-aggregate.loc[b,'macro_f1_mean']),
            'ci_2.5':float(np.percentile(delta,2.5)),'ci_97.5':float(np.percentile(delta,97.5)),'resamples':args.bootstrap})
    if pairs:pd.DataFrame(pairs).to_csv(paper/'bootstrap_differences.csv',index=False)
    (paper/'results_to_paste.md').write_text('\n\n'.join(sections),encoding='utf-8')
    images=['comparison_macro_f1.png']+(['performance_vs_latency.png'] if (paper/'performance_vs_latency.png').exists() else [])
    body=f'<h1>{html.escape(title)}</h1><p>Generated from measured runs. Review before manuscript submission.</p>'+aggregate.round(5).to_html()
    body+=''.join(f'<img src="{f}" style="max-width:100%"><p>{f}</p>' for f in images)
    for run,r in zip(run_dirs,records):
        body+=f'<h2>{html.escape(run.name)}</h2><p>Seed-specific results; do not select a seed based on test accuracy.</p>'
        for image in ['accuracy_curve.png','loss_curve.png','macro_f1_curve.png','confusion_matrix.png']:
            body+=f'<img src="../runs/{run.name}/{image}" style="width:48%">'
    (paper/'results.html').write_text('<!doctype html><meta charset="utf-8"><style>body{font-family:Arial;margin:30px}table{border-collapse:collapse;font-size:10px}td,th{padding:5px;border:1px solid #ccc}</style>'+body,encoding='utf-8')
    export_word(paper,title,aggregate,run_dirs,sections)
    print('Paper outputs:',paper.resolve())

def export_word(paper,title,aggregate,runs,sections):
    """Editable results supplement; insert reviewed sections into the full manuscript."""
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document();s=d.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.left_margin=s.right_margin=Inches(.75)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=d.styles[name];st.font.color.rgb=RGBColor(0,0,0);st.font.name='Times New Roman'
        for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    d.styles['Normal'].font.size=Pt(11)
    d.add_paragraph(title,style='Title')
    d.add_paragraph('Results supplement for the waste classification manuscript. Review measured outputs, replace the proposal wording and integrate these tables and figures with your introduction, methodology and references.')
    d.add_heading('Experimental status',1);d.add_paragraph(sections[3])
    def add_table(headers,rows):
        t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid'
        for cell,value in zip(t.rows[0].cells,headers):cell.text=value
        for row in rows:
            for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                pr.append(borders)
                if i==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');cell._tc.get_or_add_tcPr().append(sh)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
                    for run in p.runs:run.font.size=Pt(9);run.bold=i==0
        d.add_paragraph()
    def value(r,key):
        mean=r[key+'_mean'];std=r[key+'_std']
        return f'{mean:.4f} ± {std:.4f}' if pd.notna(std) else f'{mean:.4f} (one run)'
    d.add_heading('Classification performance',1)
    add_table(['Model','Accuracy','Macro precision','Macro recall','Macro F1'],
        [[m,*[value(r,k) for k in ['accuracy','macro_precision','macro_recall','macro_f1']]] for m,r in aggregate.iterrows()])
    d.add_paragraph('Metrics are fractions. Values are mean ± sample standard deviation across completed training seeds on a common fixed test split.')
    d.add_heading('Computational measurements',1)
    rows=[]
    for m,r in aggregate.iterrows():
        rows.append([m,f'{r.parameters_mean:,.0f}',f'{r.model_size_mib_mean:.2f}',f'{r.training_seconds_mean/60:.2f}',
            f'{r.model_median_ms_mean:.2f}' if 'model_median_ms_mean' in r else 'Not measured'])
    add_table(['Model','Parameters','Size MiB','Training minutes','CPU median ms'],rows)
    d.add_paragraph('Training duration includes validation macro F1 and checkpoint I/O. Inference uses float32 and batch size one. Consult environment files and latency sample CSVs for settings and variability. Uncertainty in timing and image sampling is not represented by this compact table.')
    d.add_heading('Results paragraph for review',1)
    d.add_paragraph(sections[8])
    for name in ['comparison_macro_f1.png','performance_vs_latency.png']:
        if (paper/name).exists():d.add_picture(str(paper/name),width=Inches(6.5));d.add_paragraph(name)
    for run in runs:
        if not run.name.endswith('seed42'):continue
        heading=d.add_heading(run.name.replace('_',' '),1)
        heading.paragraph_format.page_break_before=True
        d.add_paragraph('Seed 42 is a predeclared illustrative run. Use aggregate metrics for model ranking. Epoch indices start at zero; transfer curves concatenate head training and fine-tuning.')
        d.add_picture(str(run/'training_panel.png'),width=Inches(6.5))
        d.add_paragraph('Training and validation accuracy and loss')
        d.add_picture(str(run/'confusion_matrix.png'),width=Inches(4.8))
        d.add_paragraph('Confusion matrix on held-out images')
    d.add_heading('Limitations',1);d.add_paragraph(sections[-1])
    d.save(paper/'paper_results.docx')

def main():
    parser=argparse.ArgumentParser(description=__doc__);sub=parser.add_subparsers(dest='command',required=True)
    d=sub.add_parser('download');d.add_argument('--dest',default='data')
    p=sub.add_parser('prepare');p.add_argument('--data',default='data/dataset-resized');p.add_argument('--out',default='outputs')
    t=sub.add_parser('train');t.add_argument('--out',default='outputs');t.add_argument('--models',nargs='+',choices=MODELS,default=MODELS)
    t.add_argument('--seeds',nargs='+',type=int,default=[42,123,2026]);t.add_argument('--batch-size',type=int,default=32)
    t.add_argument('--epochs',type=int,default=30);t.add_argument('--head-epochs',type=int,default=10);t.add_argument('--fine-epochs',type=int,default=20)
    t.add_argument('--weights',choices=['imagenet','none'],default='imagenet');t.add_argument('--smoke',action='store_true')
    b=sub.add_parser('benchmark');b.add_argument('--out',default='outputs');b.add_argument('--calls',type=int,default=500);b.add_argument('--threads',type=int,default=1)
    r=sub.add_parser('report');r.add_argument('--out',default='outputs');r.add_argument('--bootstrap',type=int,default=1000)
    a=parser.parse_args()
    if a.command=='download':download(a.dest)
    elif a.command=='prepare':prepare(a.data,a.out)
    elif a.command=='train':
        if min(a.epochs,a.head_epochs,a.batch_size)<1 or a.fine_epochs<0:parser.error('Epoch and batch counts must be positive; fine epochs may be zero.')
        train(a)
    elif a.command=='benchmark':
        if a.calls<1 or a.threads<1:parser.error('calls and threads must be positive')
        benchmark(a)
    else:
        if a.bootstrap<1:parser.error('bootstrap must be positive')
        report(a)

if __name__=='__main__':main()
